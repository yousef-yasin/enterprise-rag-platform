"""DOCX parsing via python-docx (docs/ARCHITECTURE.md §8.2, §30.1).

python-docx reads the OOXML package with lxml, whose defaults disable network access
and external-DTD resolution and cap entity expansion (``huge_tree=False``). We add
the zip-bomb guard before opening.
"""

from __future__ import annotations

import io

from app.core.errors import UnprocessableDocumentError
from app.core.models import DocumentBlock
from app.core.parsing.normalize import normalize_text
from app.core.parsing.security import ArchiveLimits, assert_safe_zip


def parse_docx(data: bytes, limits: ArchiveLimits) -> list[DocumentBlock]:
    assert_safe_zip(data, limits)

    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(io.BytesIO(data))
    except PackageNotFoundError as exc:
        raise UnprocessableDocumentError("not a valid DOCX package", reason="corrupt") from exc
    except Exception as exc:
        raise UnprocessableDocumentError("could not read DOCX", reason="parse_error") from exc

    blocks: list[DocumentBlock] = []
    offset = 0
    for para in document.paragraphs:
        raw = normalize_text(para.text)
        if not raw:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            level = _heading_level(style)
            kind = "heading"
        elif style.startswith("list") or para.text.lstrip().startswith(("- ", "* ", "• ")):
            level, kind = 0, "list_item"
        else:
            level, kind = 0, "paragraph"
        end = offset + len(raw)
        blocks.append(
            DocumentBlock(
                kind=kind,  # type: ignore[arg-type]
                text=raw,
                level=level,
                char_start=offset,
                char_end=end,
            )
        )
        offset = end + 1

    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(normalize_text(cell.text) for cell in row.cells))
        text = "\n".join(r for r in rows if r.strip())
        if text:
            end = offset + len(text)
            blocks.append(DocumentBlock(kind="table", text=text, char_start=offset, char_end=end))
            offset = end + 1

    return blocks


def _heading_level(style: str) -> int:
    digits = "".join(c for c in style if c.isdigit())
    return int(digits) if digits else 1
