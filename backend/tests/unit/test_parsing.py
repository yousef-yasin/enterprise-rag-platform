"""Document parsing + the section-30.3 edge-case taxonomy."""

from __future__ import annotations

import io
import zipfile

import pytest

from app.core.errors import UnprocessableDocumentError, UnsupportedMediaTypeError
from app.core.parsing import ParseLimits, parse_document
from app.core.parsing.security import ArchiveLimits, assert_safe_zip


def _limits() -> ParseLimits:
    return ParseLimits(
        max_pages=1500,
        max_chunks_per_doc=5000,
        min_text_chars_per_page=30,
        min_doc_chars=20,
        archive_max_uncompressed_bytes=200 * 1024 * 1024,
        archive_max_entries=2000,
        archive_max_ratio=120,
    )


def test_parse_markdown_produces_headings_and_paragraphs() -> None:
    md = b"# Title\n\nFirst paragraph with content.\n\n## Sub\n\n- item one\n- item two\n"
    tree = parse_document(md, filename="doc.md", limits=_limits())
    kinds = [b.kind for b in tree.blocks]
    assert "heading" in kinds
    assert "paragraph" in kinds
    assert tree.title == "Title"
    assert tree.lang == "en"


def test_parse_plain_text() -> None:
    tree = parse_document(
        b"This is the first paragraph.\n\nThis is the second one, longer.",
        filename="notes.txt",
        limits=_limits(),
    )
    assert len(tree.blocks) == 2


def test_empty_file_rejected() -> None:
    with pytest.raises(UnprocessableDocumentError) as exc:
        parse_document(b"", filename="x.txt", limits=_limits())
    assert exc.value.reason == "no_text"


def test_too_little_text_rejected() -> None:
    with pytest.raises(UnprocessableDocumentError) as exc:
        parse_document(b"hi", filename="x.txt", limits=_limits())
    assert exc.value.reason == "no_text"


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedMediaTypeError):
        parse_document(b"anything at all here" * 5, filename="malware.exe", limits=_limits())


def test_pdf_magic_mismatch() -> None:
    with pytest.raises(UnsupportedMediaTypeError):
        parse_document(b"not a pdf" * 10, filename="fake.pdf", limits=_limits())


def test_docx_magic_mismatch() -> None:
    with pytest.raises(UnsupportedMediaTypeError):
        parse_document(b"not a zip" * 10, filename="fake.docx", limits=_limits())


def test_zip_bomb_rejected() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("bomb.txt", b"0" * (5 * 1024 * 1024))
    limits = ArchiveLimits(max_uncompressed_bytes=1024 * 1024, max_entries=10, max_ratio=120)
    with pytest.raises(UnprocessableDocumentError) as exc:
        assert_safe_zip(buf.getvalue(), limits)
    assert exc.value.reason == "archive_limits"


def test_zip_bomb_entry_count_rejected() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(50):
            zf.writestr(f"f{i}.txt", b"x")
    limits = ArchiveLimits(max_uncompressed_bytes=10**9, max_entries=10, max_ratio=10**9)
    with pytest.raises(UnprocessableDocumentError):
        assert_safe_zip(buf.getvalue(), limits)


def test_real_docx_parses() -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Company Handbook", level=1)
    doc.add_paragraph("Welcome to the team. This paragraph has enough text to pass the check.")
    doc.add_paragraph("A second paragraph with more useful content for chunking purposes.")
    buf = io.BytesIO()
    doc.save(buf)

    tree = parse_document(buf.getvalue(), filename="handbook.docx", limits=_limits())
    assert tree.title == "Company Handbook"
    assert any(b.kind == "heading" for b in tree.blocks)
    assert tree.total_chars > 20


def test_real_pdf_parses() -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for i in range(8):
        pdf.cell(
            0,
            8,
            text=f"This is line {i} of the sample PDF document used for testing.",
            new_x="LMARGIN",
            new_y="NEXT",
        )
    data = bytes(pdf.output())

    tree = parse_document(data, filename="sample.pdf", limits=_limits())
    assert tree.page_count == 1
    assert tree.total_chars > 30
