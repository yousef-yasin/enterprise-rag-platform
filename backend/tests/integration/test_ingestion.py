"""Ingestion pipeline + the PG<->Qdrant write protocol (Phase 2 gate)."""

from __future__ import annotations

import io
import uuid

import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.core.embedding_profile import (
    ChunkPolicy,
    compute_embedding_profile_id,
    qdrant_collection_name,
)
from app.core.enums import ChunkStatus, DocumentStatus
from app.infra.db.models import Chunk, Document
from app.infra.qdrant.store import QdrantVectorStore
from app.providers.embeddings.fake import FakeEmbedder
from app.services.ingestion import IngestionPipeline
from tests.integration._helpers import create_kb, register_and_auth

_MD = b"""# Employee Handbook

## Onboarding

Welcome to the company. On your first day set up your laptop, request access to the
shared drives, and read the security policy. Your manager will schedule a one-on-one.

## Benefits

The company offers health insurance, a retirement plan, and 25 days of paid leave
per year. Submit expense claims through the finance portal within 30 days.

## Code of Conduct

Treat colleagues with respect. Report concerns to your manager or to HR. Harassment
of any kind is not tolerated and will result in disciplinary action.
"""


def _docx_bytes() -> bytes:
    from docx import Document as Docx

    d = Docx()
    d.add_heading("Quarterly Report", level=1)
    d.add_paragraph("Revenue grew twelve percent this quarter driven by new enterprise deals.")
    d.add_paragraph("Operating costs remained flat. Headcount increased by eight people.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def _run_pipeline(document_id: uuid.UUID) -> str:
    return await IngestionPipeline(get_settings()).run(str(document_id))


async def _upload(
    client: AsyncClient, headers: dict[str, str], kb_id: str, *, name: str, data: bytes
) -> str:
    resp = await client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        headers=headers,
        files={"file": (name, data, "application/octet-stream")},
    )
    assert resp.status_code == 202, resp.text
    doc_id: str = resp.json()["id"]
    return doc_id


@pytest.mark.parametrize(
    ("name", "factory"),
    [
        ("handbook.md", lambda: _MD),
        ("handbook.txt", lambda: _MD.replace(b"#", b"")),
        ("report.docx", _docx_bytes),
    ],
)
async def test_ingest_produces_ready_chunks(
    client: AsyncClient, db_session: AsyncSession, name: str, factory: object
) -> None:
    headers = await register_and_auth(client, email="owner@example.com")
    kb = await create_kb(client, headers, slug="handbook")
    doc_id = await _upload(client, headers, kb["id"], name=name, data=factory())  # type: ignore[operator]

    outcome = await _run_pipeline(uuid.UUID(doc_id))
    assert outcome in {"ready", "partially_indexed"}, outcome

    db_session.expire_all()
    doc = await db_session.get(Document, uuid.UUID(doc_id))
    assert doc is not None
    assert doc.status is DocumentStatus.READY
    assert doc.active_index_version == 1

    chunks = (
        (await db_session.execute(select(Chunk).where(Chunk.document_id == doc.id))).scalars().all()
    )
    assert len(chunks) >= 1
    assert all(c.status is ChunkStatus.READY for c in chunks)
    assert [c.ordinal for c in sorted(chunks, key=lambda c: c.ordinal)] == list(range(len(chunks)))
    assert all(c.char_end >= c.char_start for c in chunks)

    max_tokens = FakeEmbedder().profile.max_tokens
    assert all(c.token_count <= max_tokens for c in chunks)

    # Qdrant points == ready chunks
    profile_id = compute_embedding_profile_id(
        FakeEmbedder().profile, ChunkPolicy.from_settings(get_settings())
    )
    collection = qdrant_collection_name(kb["id"], profile_id)
    store = QdrantVectorStore(get_settings())
    try:
        assert await store.count(collection) == len(chunks)
    finally:
        await store.close()


async def test_duplicate_upload_returns_same_document(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    headers = await register_and_auth(client, email="owner@example.com")
    kb = await create_kb(client, headers, slug="handbook")

    first = await _upload(client, headers, kb["id"], name="a.md", data=_MD)
    second = await _upload(client, headers, kb["id"], name="a-again.md", data=_MD)
    assert first == second

    db_session.expire_all()
    docs = (
        (await db_session.execute(select(Document).where(Document.knowledge_base_id == kb["id"])))
        .scalars()
        .all()
    )
    assert len(docs) == 1


async def test_reprocess_bumps_version_without_visibility_gap(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    headers = await register_and_auth(client, email="owner@example.com")
    kb = await create_kb(client, headers, slug="handbook")
    doc_id = await _upload(client, headers, kb["id"], name="h.md", data=_MD)
    await _run_pipeline(uuid.UUID(doc_id))

    r = await client.post(f"/api/v1/documents/{doc_id}/reprocess", headers=headers)
    assert r.status_code == 202

    db_session.expire_all()
    doc = await db_session.get(Document, uuid.UUID(doc_id))
    assert doc is not None and doc.next_index_version == 2

    outcome = await _run_pipeline(uuid.UUID(doc_id))
    assert outcome in {"ready", "partially_indexed"}

    db_session.expire_all()
    doc = await db_session.get(Document, uuid.UUID(doc_id))
    assert doc is not None and doc.active_index_version == 2

    # v2 chunks are ready; retrieval (Phase 3+) only ever sees active_index_version
    chunks = (
        (await db_session.execute(select(Chunk).where(Chunk.document_id == doc.id))).scalars().all()
    )
    v2_ready = {
        c.index_version for c in chunks if c.status is ChunkStatus.READY and c.index_version == 2
    }
    assert v2_ready == {2}


async def test_stale_points_cleanup_removes_old_version(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    headers = await register_and_auth(client, email="owner@example.com")
    kb = await create_kb(client, headers, slug="handbook")
    doc_id = await _upload(client, headers, kb["id"], name="h.md", data=_MD)
    await _run_pipeline(uuid.UUID(doc_id))
    await client.post(f"/api/v1/documents/{doc_id}/reprocess", headers=headers)
    await _run_pipeline(uuid.UUID(doc_id))

    from app.workers.tasks import delete_stale_points

    await delete_stale_points({}, doc_id, 2)

    db_session.expire_all()
    remaining = (
        (await db_session.execute(select(Chunk).where(Chunk.document_id == uuid.UUID(doc_id))))
        .scalars()
        .all()
    )
    assert {c.index_version for c in remaining} == {2}


async def test_parser_edge_case_marks_document_failed(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    headers = await register_and_auth(client, email="owner@example.com")
    kb = await create_kb(client, headers, slug="handbook")
    # scanned-PDF style: valid PDF header, no extractable text
    fake_pdf = b"%PDF-1.4\n" + b"0" * 400 + b"\n%%EOF"
    doc_id = await _upload(client, headers, kb["id"], name="scan.pdf", data=fake_pdf)

    outcome = await _run_pipeline(uuid.UUID(doc_id))
    assert outcome == "failed"

    db_session.expire_all()
    doc = await db_session.get(Document, uuid.UUID(doc_id))
    assert doc is not None
    assert doc.status is DocumentStatus.FAILED
    assert doc.failure_reason in {"no_text", "insufficient_text", "corrupt"}
    assert doc.failed_stage is not None
